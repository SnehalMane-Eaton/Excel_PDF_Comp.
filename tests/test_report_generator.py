from pathlib import Path

from openpyxl import load_workbook

from app.reporting.report_generator import generate_report


def test_generate_report_preserves_source_rows_and_pair_isolation(tmp_path):
    batch_result = {
        "summary": {
            "total_pairs": 2,
            "verified": 2,
            "processing_error": 0,
            "excel_only": 0,
            "pdf_only": 0,
            "match": 1,
            "qty_mismatch": 1,
            "bom_excel_only": 0,
            "bom_pdf_only": 0,
            "bom_review": 1,
        },
        "pairs": [
            {
                "key": "job_alpha",
                "status": "VERIFIED",
                "message": "ok",
                "excel_file": "C:/local/job_alpha.xlsx",
                "pdf_file": "C:/local/job_alpha.pdf",
                "warnings": [],
                "excel_records": [
                    {"part_number": "001", "quantity": 1, "description": "Excel alpha", "sheet_name": "Main", "excel_row": 2},
                    {"part_number": "", "quantity": 0, "description": "Blank review", "sheet_name": "Main", "excel_row": 3},
                ],
                "pdf_records": [
                    {"part_number": "001", "quantity": 1, "description": "PDF alpha", "pdf_page": 1},
                ],
                "verification": {
                    "results": [
                        {"part_number": "001", "description": "Excel alpha", "excel_quantity": 1, "pdf_quantity": 1, "difference": 0, "status": "MATCH", "excel_rows": [2], "pdf_pages": [1], "remark": ""},
                        {"part_number": "", "description": "Blank review", "excel_quantity": 0, "pdf_quantity": None, "difference": None, "status": "REVIEW", "excel_rows": [3], "pdf_pages": [], "remark": "Blank Part Number in Excel"},
                    ],
                    "summary": {"total": 2, "match": 1, "qty_mismatch": 0, "excel_only": 0, "pdf_only": 0, "review": 1},
                },
            },
            {
                "key": "job_beta",
                "status": "VERIFIED",
                "message": "ok",
                "excel_file": "C:/local/job_beta.xlsx",
                "pdf_file": "C:/local/job_beta.pdf",
                "warnings": [],
                "excel_records": [
                    {"part_number": "001", "quantity": 3, "description": "Excel beta", "sheet_name": "Main", "excel_row": 5},
                ],
                "pdf_records": [
                    {"part_number": "001", "quantity": 9, "description": "PDF beta", "pdf_page": 2},
                ],
                "verification": {
                    "results": [
                        {"part_number": "001", "description": "Excel beta", "excel_quantity": 3, "pdf_quantity": 9, "difference": 6, "status": "QTY_MISMATCH", "excel_rows": [5], "pdf_pages": [2], "remark": "Part number matched but quantity differs"},
                    ],
                    "summary": {"total": 1, "match": 0, "qty_mismatch": 1, "excel_only": 0, "pdf_only": 0, "review": 0},
                },
            },
        ],
    }

    output = tmp_path / "report.xlsx"
    generate_report(batch_result, output)

    workbook = load_workbook(output)
    assert workbook.sheetnames[:2] == ["SUMMARY", "VERIFICATION_AUDIT"]
    first_sheet = workbook["01_job_alpha"]
    second_sheet = workbook["02_job_beta"]

    assert [first_sheet.cell(1, column).value for column in range(1, 10)] == [
        "BOM1-PART",
        "BOM1-QTY",
        "BOM1-DESCRIPTION",
        "BOM1-RESULT",
        "QTY RESULT",
        "BOM2-PART",
        "BOM2-QTY",
        "BOM2-DESCRIPTION",
        "BOM2-RESULT",
    ]
    assert first_sheet["A2"].value == "001"
    assert first_sheet["C2"].value == "Excel alpha"
    assert first_sheet["F2"].value == "001"
    assert first_sheet["H2"].value == "PDF alpha"
    assert first_sheet["C3"].value == "Blank review"
    assert "_PAIR_001_AUDIT" in first_sheet["D2"].value
    assert "_PAIR_002_AUDIT" in second_sheet["D2"].value
    assert "$1000" not in first_sheet["D2"].value

    audit_sheet = workbook["VERIFICATION_AUDIT"]
    audit_rows = list(audit_sheet.iter_rows(min_row=2, max_col=10, values_only=True))
    assert ("job_alpha", None, "REVIEW", 0, None, None, "Blank review", "3", None, "Blank Part Number in Excel") in audit_rows


def test_generate_report_swaps_display_order_and_preserves_text_literals(tmp_path):
    batch_result = {
        "summary": {"total_pairs": 1, "verified": 1, "processing_error": 0, "excel_only": 0, "pdf_only": 0, "match": 1, "qty_mismatch": 0, "bom_excel_only": 0, "bom_pdf_only": 0, "bom_review": 0},
        "pairs": [
            {
                "key": "literal_case",
                "status": "VERIFIED",
                "message": "ok",
                "excel_file": "excel.xlsx",
                "pdf_file": "pdf.pdf",
                "warnings": [],
                "excel_records": [{"part_number": "=LEAD*01", "quantity": "0", "description": "=desc", "sheet_name": "Main", "excel_row": 2}],
                "pdf_records": [{"part_number": "001", "quantity": 0, "description": "PDF desc", "pdf_page": 1}],
                "verification": {"results": [{"part_number": "001", "description": "PDF desc", "excel_quantity": 0, "pdf_quantity": 0, "difference": 0, "status": "MATCH", "excel_rows": [2], "pdf_pages": [1], "remark": ""}], "summary": {"total": 1, "match": 1, "qty_mismatch": 0, "excel_only": 0, "pdf_only": 0, "review": 0}},
            }
        ],
    }

    output = tmp_path / "swapped.xlsx"
    generate_report(batch_result, output, bom1_type="PDF", bom2_type="Excel")
    workbook = load_workbook(output)
    sheet = workbook["01_literal_case"]

    assert sheet["A2"].value == "001"
    excel_parts = [sheet[f"F{row}"].value for row in range(2, sheet.max_row + 1)]
    excel_descriptions = [sheet[f"H{row}"].value for row in range(2, sheet.max_row + 1)]
    literal_row = excel_parts.index("=LEAD*01") + 2
    assert sheet[f"F{literal_row}"].data_type == "s"
    assert excel_descriptions[literal_row - 2] == "=desc"
    assert sheet[f"H{literal_row}"].data_type == "s"
    assert sheet["D2"].value.startswith("=")
    assert sheet["I2"].value.startswith("=")
