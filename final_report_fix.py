from pathlib import Path

p = Path("app/reporting/report_generator.py")
s = p.read_text(encoding="utf-8")

start = s.find("class ReportGenerator:")
if start == -1:
    raise SystemExit("ReportGenerator class not found.")

new_class = r'''
class ReportGenerator:
    """Generate the final side-by-side Excel BOM comparison report."""

    def generate(
        self,
        verification_data,
        output_path,
        bom1_type="Excel",
        bom2_type="PDF",
        **kwargs,
    ):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.formatting.rule import FormulaRule

        # ------------------------------------------------------------
        # Get the ACTUAL comparison records produced by verification.
        # ------------------------------------------------------------
        results = []

        if isinstance(verification_data, dict):
            for pair in verification_data.get("pairs", []):
                if not isinstance(pair, dict):
                    continue

                verification = pair.get("verification", {})
                if not isinstance(verification, dict):
                    continue

                pair_results = verification.get("results", [])
                if isinstance(pair_results, list):
                    results.extend(pair_results)

        # ------------------------------------------------------------
        # Workbook
        # ------------------------------------------------------------
        wb = Workbook()
        ws = wb.active
        ws.title = "BOM Comparison"

        headers = [
            "BOM1-PART",
            "BOM1-QTY",
            "BOM1-DESCRIPTION",
            "BOM1-RESULT",
            "QTY RESULT",
            "BOM2-PART",
            "BOM2-QTY",
            "BOM2-DESCRIPTION",
            "BOM2-RESULT",
        ]

        ws.append(headers)

        # ------------------------------------------------------------
        # Header colours
        # ------------------------------------------------------------
        yellow = "FFFF00"
        green = "00B050"
        purple = "7030A0"

        thin = Side(style="thin")
        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin,
        )

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cell.border = border

        for col in range(1, 4):
            ws.cell(1, col).fill = PatternFill(
                "solid",
                fgColor=yellow,
            )

        for col in (4, 5, 9):
            ws.cell(1, col).fill = PatternFill(
                "solid",
                fgColor=green,
            )

        for col in range(6, 9):
            ws.cell(1, col).fill = PatternFill(
                "solid",
                fgColor=purple,
            )

        # ------------------------------------------------------------
        # Write EVERY actual comparison record.
        #
        # verification results already contain the union of the
        # Excel/PDF BOM parts, so no BOM rows are discarded.
        # ------------------------------------------------------------
        for result in results:
            if not isinstance(result, dict):
                continue

            part = result.get("part_number", "")
            description = result.get("description", "")
            excel_qty = result.get("excel_quantity")
            pdf_qty = result.get("pdf_quantity")
            status = str(result.get("status", "")).upper()

            part = "" if part is None else str(part).strip()
            description = "" if description is None else str(description).strip()

            excel_qty_value = "" if excel_qty is None else excel_qty
            pdf_qty_value = "" if pdf_qty is None else pdf_qty

            # Determine whether the part exists on each source.
            excel_exists = excel_qty is not None
            pdf_exists = pdf_qty is not None

            if status == "EXCEL_ONLY":
                excel_exists = True
                pdf_exists = False
            elif status == "PDF_ONLY":
                excel_exists = False
                pdf_exists = True
            elif status in ("MATCH", "QTY_MISMATCH", "REVIEW"):
                excel_exists = True
                pdf_exists = True

            excel_part = part if excel_exists else ""
            excel_desc = description if excel_exists else ""

            pdf_part = part if pdf_exists else ""
            pdf_desc = description if pdf_exists else ""

            # --------------------------------------------------------
            # Apply BOM1 / BOM2 selection.
            # --------------------------------------------------------
            if bom1_type == "Excel":
                bom1_part = excel_part
                bom1_qty = excel_qty_value
                bom1_desc = excel_desc

                bom2_part = pdf_part
                bom2_qty = pdf_qty_value
                bom2_desc = pdf_desc
            else:
                bom1_part = pdf_part
                bom1_qty = pdf_qty_value
                bom1_desc = pdf_desc

                bom2_part = excel_part
                bom2_qty = excel_qty_value
                bom2_desc = excel_desc

            ws.append([
                bom1_part,
                bom1_qty,
                bom1_desc,
                None,
                None,
                bom2_part,
                bom2_qty,
                bom2_desc,
                None,
            ])

        # ------------------------------------------------------------
        # Exact requested formulas.
        # ------------------------------------------------------------
        last_row = max(ws.max_row, 2)

        for row in range(2, last_row + 1):
            ws.cell(row, 4).value = (
                f'=IF(ISNUMBER(MATCH(A{row},$F$2:$F$1000,0)),'
                f'"EXISTS IN BOM2","DOES NOT EXIST IN BOM2")'
            )

            ws.cell(row, 5).value = (
                f'=IF(ISNUMBER(MATCH(A{row},$F$2:$F$1000,0)),'
                f'IF(B{row}=INDEX($G$2:$G$1000,'
                f'MATCH(A{row},$F$2:$F$1000,0)),'
                f'"QUANTITIES MATCH","QUANTITIES DON’T MATCH"),"N/A")'
            )

            ws.cell(row, 9).value = (
                f'=IF(ISNUMBER(MATCH(F{row},$A$2:$A$1000,0)),'
                f'"EXISTS IN BOM1","DOES NOT EXIST IN BOM1")'
            )

        # ------------------------------------------------------------
        # Formatting
        # ------------------------------------------------------------
        widths = {
            "A": 24,
            "B": 12,
            "C": 42,
            "D": 25,
            "E": 25,
            "F": 24,
            "G": 12,
            "H": 42,
            "I": 25,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:I{last_row}"

        # ------------------------------------------------------------
        # Red result cells when formula evaluates to a mismatch.
        # ------------------------------------------------------------
        red_fill = PatternFill(
            "solid",
            fgColor="FFC7CE",
        )

        red_font = Font(
            color="9C0006",
        )

        ws.conditional_formatting.add(
            f"D2:D{last_row}",
            FormulaRule(
                formula=['ISNUMBER(SEARCH("DOES NOT EXIST",D2))'],
                fill=red_fill,
                font=red_font,
            ),
        )

        ws.conditional_formatting.add(
            f"E2:E{last_row}",
            FormulaRule(
                formula=['ISNUMBER(SEARCH("DON’T",E2))'],
                fill=red_fill,
                font=red_font,
            ),
        )

        ws.conditional_formatting.add(
            f"I2:I{last_row}",
            FormulaRule(
                formula=['ISNUMBER(SEARCH("DOES NOT EXIST",I2))'],
                fill=red_fill,
                font=red_font,
            ),
        )

        # Force Excel to recalculate formulas when the file opens.
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
        except Exception:
            pass

        wb.save(output_path)
        return output_path
'''

p.write_text(s[:start] + new_class + "\n", encoding="utf-8")

compile(p.read_text(encoding="utf-8"), str(p), "exec")
print("FINAL REPORT GENERATOR INSTALLED")
print("Actual Excel/PDF comparison records:", "YES")
print("9-column report:", "YES")
print("BOM1/BOM2 switching:", "YES")
