from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


# ============================================================
# REPORT CONFIGURATION
# ============================================================

REPORT_SHEET_NAME = "BOM COMPARISON"

# User-provided formula ranges must remain exactly like this.
FORMULA_MAX_ROW = 1000

HEADERS = [
    "BOM1-PART",
    "BOM1-QTY",
    "BOM1-DESCRIPTION",
    "BOM1-RESULT",
    "QTY RESULT",
    "BOM2-PART",
    "BOM2-QTY",
    "BOM2-DESCRIPTION",
    "BOM2-RESULT",
]

# Approximate colors from the user's reference workbook.
HEADER_YELLOW = "FFFF00"
HEADER_GREEN = "00B050"
HEADER_PURPLE = "7030A0"

WHITE = "FFFFFF"
BLACK = "000000"
RED = "FF0000"

# Result-cell fills.
RESULT_GREEN = "C6EFCE"
RESULT_RED = "FFC7CE"
RESULT_YELLOW = "FFEB9C"

THIN_BLACK = Side(
    style="thin",
    color=BLACK,
)

CELL_BORDER = Border(
    left=THIN_BLACK,
    right=THIN_BLACK,
    top=THIN_BLACK,
    bottom=THIN_BLACK,
)


# ============================================================
# PUBLIC API
# ============================================================

def generate_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    """
    Generate the BOM comparison report.

    Parameters
    ----------
    verification_data:
        Verification/extraction data.

    output_path:
        Destination .xlsx path.

    bom1_type:
        "Excel" or "PDF".

    bom2_type:
        "Excel" or "PDF".

    Returns
    -------
    str
        Saved report path.

    Notes
    -----
    BOM1 and BOM2 are intentionally independent lists.

    Example:

        BOM1 = Excel
        BOM2 = PDF

    produces:

        Excel parts -> columns A:C
        PDF parts   -> columns F:H

    If the user selects:

        BOM1 = PDF
        BOM2 = Excel

    the data is automatically reversed.
    """

    bom1_type = _normalize_bom_type(bom1_type)
    bom2_type = _normalize_bom_type(bom2_type)

    if bom1_type == bom2_type:
        raise ValueError(
            "BOM1 and BOM2 must be different. "
            "Select Excel and PDF."
        )

    excel_records = _extract_excel_records(verification_data)
    pdf_records = _extract_pdf_records(verification_data)

    if bom1_type == "Excel":
        bom1_records = excel_records
        bom2_records = pdf_records
    else:
        bom1_records = pdf_records
        bom2_records = excel_records

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = REPORT_SHEET_NAME

    _configure_worksheet(worksheet)

    _write_headers(worksheet)

    max_rows = max(
        len(bom1_records),
        len(bom2_records),
        1,
    )

    for index in range(max_rows):
        excel_row = index + 2

        bom1_record = (
            bom1_records[index]
            if index < len(bom1_records)
            else {}
        )

        bom2_record = (
            bom2_records[index]
            if index < len(bom2_records)
            else {}
        )

        _write_bom1_row(
            worksheet=worksheet,
            row=excel_row,
            record=bom1_record,
        )

        _write_bom2_row(
            worksheet=worksheet,
            row=excel_row,
            record=bom2_record,
        )

        _write_result_formulas(
            worksheet=worksheet,
            row=excel_row,
        )

    _apply_result_conditional_formatting(
        worksheet=worksheet,
        first_data_row=2,
        last_data_row=max_rows + 1,
    )

    _apply_row_formatting(
        worksheet=worksheet,
        first_data_row=2,
        last_data_row=max_rows + 1,
    )

    _freeze_and_filter(
        worksheet=worksheet,
        last_row=max_rows + 1,
    )

    _configure_print_settings(
        worksheet=worksheet,
        last_row=max_rows + 1,
    )

    # Make Excel recalculate formulas when the workbook is opened.
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    output = Path(output_path)
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(output)

    return str(output)


# ============================================================
# BOM TYPE
# ============================================================

def _normalize_bom_type(value: Any) -> str:
    text = str(value or "").strip().lower()

    if text == "excel":
        return "Excel"

    if text == "pdf":
        return "PDF"

    raise ValueError(
        f"Invalid BOM type: {value!r}. "
        "Expected 'Excel' or 'PDF'."
    )


# ============================================================
# DATA EXTRACTION FROM VERIFICATION RESULT
# ============================================================

def _extract_excel_records(
    verification_data: dict[str, Any],
) -> list[dict[str, Any]]:
    """
    Build an independent Excel BOM list.

    The function supports the existing verification structure:

        {
            "results": [
                {
                    "part_number": ...,
                    "excel_quantity": ...,
                    "description": ...
                }
            ]
        }

    It also supports direct:

        {
            "excel_records": [...]
        }

    This keeps the report generator compatible with the
    existing application architecture.
    """

    direct_records = verification_data.get("excel_records")

    if direct_records is not None:
        return _normalize_records(
            direct_records,
            source="excel",
        )

    results = verification_data.get("results", [])

    records: list[dict[str, Any]] = []

    for result in results:
        if not isinstance(result, dict):
            continue

        part_number = result.get("part_number")

        if not _has_part_number(part_number):
            continue

        records.append(
            {
                "part_number": part_number,
                "quantity": result.get("excel_quantity"),
                "description": result.get("description", ""),
            }
        )

    return records


def _extract_pdf_records(
    verification_data: dict[str, Any],
) -> list[dict[str, Any]]:
    """
    Build an independent PDF BOM list.

    Supports:

        verification_data["pdf_records"]

    or the existing verification result structure.
    """

    direct_records = verification_data.get("pdf_records")

    if direct_records is not None:
        return _normalize_records(
            direct_records,
            source="pdf",
        )

    results = verification_data.get("results", [])

    records: list[dict[str, Any]] = []

    for result in results:
        if not isinstance(result, dict):
            continue

        part_number = result.get("part_number")

        if not _has_part_number(part_number):
            continue

        records.append(
            {
                "part_number": part_number,
                "quantity": result.get("pdf_quantity"),
                "description": result.get("description", ""),
            }
        )

    return records


def _normalize_records(
    records: Iterable[Any],
    source: str,
) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []

    for record in records:
        if not isinstance(record, dict):
            continue

        if source == "excel":
            part_number = (
                record.get("part_number")
                or record.get("Part Number")
                or record.get("PART NUMBER")
                or record.get("BOM1-PART")
            )

            quantity = (
                record.get("quantity")
                if "quantity" in record
                else record.get("QTY")
            )

        else:
            part_number = (
                record.get("part_number")
                or record.get("Part Number")
                or record.get("PART NUMBER")
                or record.get("BOM2-PART")
            )

            quantity = (
                record.get("quantity")
                if "quantity" in record
                else record.get("QTY")
            )

        description = (
            record.get("description")
            or record.get("Description")
            or record.get("DESCRIPTION")
            or ""
        )

        if not _has_part_number(part_number):
            continue

        normalized.append(
            {
                "part_number": str(part_number).strip(),
                "quantity": quantity,
                "description": str(description).strip(),
            }
        )

    return normalized


def _has_part_number(value: Any) -> bool:
    if value is None:
        return False

    text = str(value).strip()

    if not text:
        return False

    return True


# ============================================================
# WORKSHEET CONFIGURATION
# ============================================================

def _configure_worksheet(worksheet) -> None:
    worksheet.sheet_view.showGridLines = False

    worksheet.freeze_panes = "A2"

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.sheet_view.zoomScale = 90


def _write_headers(worksheet) -> None:
    for column_index, header in enumerate(
        HEADERS,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color=BLACK,
            size=11,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = CELL_BORDER

        if column_index in (1, 2, 3):
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_YELLOW,
            )

        elif column_index in (4, 5, 9):
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_GREEN,
            )

        elif column_index in (6, 7, 8):
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_PURPLE,
            )

            cell.font = Font(
                bold=True,
                color=WHITE,
                size=11,
            )

    worksheet.row_dimensions[1].height = 30


# ============================================================
# DATA ROWS
# ============================================================

def _write_bom1_row(
    worksheet,
    row: int,
    record: dict[str, Any],
) -> None:
    worksheet.cell(
        row=row,
        column=1,
        value=_clean_value(record.get("part_number")),
    )

    worksheet.cell(
        row=row,
        column=2,
        value=_clean_quantity(record.get("quantity")),
    )

    worksheet.cell(
        row=row,
        column=3,
        value=_clean_value(record.get("description")),
    )


def _write_bom2_row(
    worksheet,
    row: int,
    record: dict[str, Any],
) -> None:
    worksheet.cell(
        row=row,
        column=6,
        value=_clean_value(record.get("part_number")),
    )

    worksheet.cell(
        row=row,
        column=7,
        value=_clean_quantity(record.get("quantity")),
    )

    worksheet.cell(
        row=row,
        column=8,
        value=_clean_value(record.get("description")),
    )


def _clean_value(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _clean_quantity(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(value, bool):
        return int(value)

    try:
        number = float(value)

        if number.is_integer():
            return int(number)

        return number

    except (ValueError, TypeError):
        return value


# ============================================================
# EXACT FORMULAS
# ============================================================

def _write_result_formulas(
    worksheet,
    row: int,
) -> None:
    """
    Write the user's formulas.

    D = BOM1 RESULT
    E = QTY RESULT
    I = BOM2 RESULT

    The user's formulas are preserved logically and the
    current row is substituted automatically.
    """

    # --------------------------------------------------------
    # BOM1 RESULT
    #
    # User formula:
    #
    # =IF(ISNUMBER(MATCH(A2,$F$2:$F$1000,0)),
    # "EXISTS IN BOM2",
    # "DOES NOT EXIST IN BOM2")
    # --------------------------------------------------------

    worksheet.cell(
        row=row,
        column=4,
        value=(
            f'=IF(ISNUMBER(MATCH(A{row},'
            f'$F$2:$F$1000,0)),'
            f'"EXISTS IN BOM2",'
            f'"DOES NOT EXIST IN BOM2")'
        ),
    )

    # --------------------------------------------------------
    # QTY RESULT
    #
    # User formula:
    #
    # =IF(ISNUMBER(MATCH(A5,$F$2:$F$1000,0)),
    # IF(B5=INDEX($G$2:$G$1000,
    # MATCH(A5,$F$2:$F$1000,0)),
    # "QUANTITIES MATCH",
    # "QUANTITIES DON’T MATCH"),
    # "N/A")
    # --------------------------------------------------------

    worksheet.cell(
        row=row,
        column=5,
        value=(
            f'=IF(ISNUMBER(MATCH(A{row},'
            f'$F$2:$F$1000,0)),'
            f'IF(B{row}=INDEX($G$2:$G$1000,'
            f'MATCH(A{row},$F$2:$F$1000,0)),'
            f'"QUANTITIES MATCH",'
            f'"QUANTITIES DON’T MATCH"),'
            f'"N/A")'
        ),
    )

    # --------------------------------------------------------
    # BOM2 RESULT
    #
    # User formula:
    #
    # =IF(ISNUMBER(MATCH(F4,$A$2:$A$1000,0)),
    # "EXISTS IN BOM1",
    # "DOES NOT EXIST IN BOM1")
    # --------------------------------------------------------

    worksheet.cell(
        row=row,
        column=9,
        value=(
            f'=IF(ISNUMBER(MATCH(F{row},'
            f'$A$2:$A$1000,0)),'
            f'"EXISTS IN BOM1",'
            f'"DOES NOT EXIST IN BOM1")'
        ),
    )


# ============================================================
# RESULT CONDITIONAL FORMATTING
# ============================================================

def _apply_result_conditional_formatting(
    worksheet,
    first_data_row: int,
    last_data_row: int,
) -> None:
    """
    Apply visual formatting to the result columns.

    D:
        EXISTS IN BOM2 -> green
        DOES NOT EXIST IN BOM2 -> red

    E:
        QUANTITIES MATCH -> green
        QUANTITIES DON'T MATCH -> red
        N/A -> yellow

    I:
        EXISTS IN BOM1 -> green
        DOES NOT EXIST IN BOM1 -> red
    """

    # --------------------------------------------------------
    # BOM1 RESULT - D
    # --------------------------------------------------------

    d_range = (
        f"D{first_data_row}:D{last_data_row}"
    )

    worksheet.conditional_formatting.add(
        d_range,
        FormulaRule(
            formula=[
                f'D{first_data_row}="EXISTS IN BOM2"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_GREEN,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        d_range,
        FormulaRule(
            formula=[
                f'D{first_data_row}="DOES NOT EXIST IN BOM2"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_RED,
            ),
            font=Font(
                color=BLACK,
            ),
        ),
    )

    # --------------------------------------------------------
    # QTY RESULT - E
    # --------------------------------------------------------

    e_range = (
        f"E{first_data_row}:E{last_data_row}"
    )

    worksheet.conditional_formatting.add(
        e_range,
        FormulaRule(
            formula=[
                f'E{first_data_row}="QUANTITIES MATCH"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_GREEN,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        e_range,
        FormulaRule(
            formula=[
                f'E{first_data_row}="QUANTITIES DON’T MATCH"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_RED,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        e_range,
        FormulaRule(
            formula=[
                f'E{first_data_row}="N/A"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_YELLOW,
            ),
        ),
    )

    # --------------------------------------------------------
    # BOM2 RESULT - I
    # --------------------------------------------------------

    i_range = (
        f"I{first_data_row}:I{last_data_row}"
    )

    worksheet.conditional_formatting.add(
        i_range,
        FormulaRule(
            formula=[
                f'I{first_data_row}="EXISTS IN BOM1"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_GREEN,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        i_range,
        FormulaRule(
            formula=[
                f'I{first_data_row}="DOES NOT EXIST IN BOM1"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=RESULT_RED,
            ),
            font=Font(
                color=BLACK,
            ),
        ),
    )


# ============================================================
# CELL FORMATTING
# ============================================================

def _apply_row_formatting(
    worksheet,
    first_data_row: int,
    last_data_row: int,
) -> None:
    for row in range(
        first_data_row,
        last_data_row + 1,
    ):
        for column in range(1, 10):
            cell = worksheet.cell(
                row=row,
                column=column,
            )

            cell.border = CELL_BORDER

            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "center"
                    if column in (2, 4, 5, 7, 9)
                    else "left"
                ),
                wrap_text=True,
            )

            cell.font = Font(
                name="Calibri",
                size=11,
                color=BLACK,
            )

        worksheet.row_dimensions[row].height = 22


# ============================================================
# COLUMN WIDTHS
# ============================================================

def _freeze_and_filter(
    worksheet,
    last_row: int,
) -> None:
    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        f"A1:I{last_row}"
    )

    widths = {
        "A": 20,
        "B": 12,
        "C": 38,
        "D": 28,
        "E": 28,
        "F": 20,
        "G": 12,
        "H": 38,
        "I": 28,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


# ============================================================
# PRINT SETTINGS
# ============================================================

def _configure_print_settings(
    worksheet,
    last_row: int,
) -> None:
    worksheet.print_area = (
        f"A1:I{last_row}"
    )

    worksheet.page_setup.orientation = "landscape"

    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_title_rows = "1:1"

    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    worksheet.oddFooter.center.text = (
        "BOM Comparison Report"
    )

    worksheet.oddFooter.right.text = (
        "Page &[Page] of &[Pages]"
    )


# ============================================================
# OPTIONAL COMPATIBILITY WRAPPERS
# ============================================================

def create_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    """
    Compatibility alias for code that calls create_report().
    """
    return generate_report(
        verification_data=verification_data,
        output_path=output_path,
        bom1_type=bom1_type,
        bom2_type=bom2_type,
    )


def save_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    """
    Compatibility alias for code that calls save_report().
    """
    return generate_report(
        verification_data=verification_data,
        output_path=output_path,
        bom1_type=bom1_type,
        bom2_type=bom2_type,
    )
# =========================================================
# Compatibility class used by MainWindow
# =========================================================

class ReportGenerator:
    """
    UI compatibility wrapper.

    The verification engine returns:
        batch_result["pairs"][...]["verification"]["results"]

    The report is generated directly from those actual comparison
    results so the selected Excel/PDF BOM data appears in Excel.
    """

    def generate(
        self,
        verification_data,
        output_path,
        bom1_type="Excel",
        bom2_type="PDF",
        **kwargs,
    ):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.utils import get_column_letter

        # ---------------------------------------------------------
        # Collect actual verification results from BatchProcessor
        # ---------------------------------------------------------
        results = []

        if isinstance(verification_data, dict):
            for pair in verification_data.get("pairs", []):
                if not isinstance(pair, dict):
                    continue

                verification = pair.get("verification", {})
                if isinstance(verification, dict):
                    rows = verification.get("results", [])
                    if isinstance(rows, list):
                        results.extend(rows)

        # ---------------------------------------------------------
        # Helpers for different result dictionary structures
        # ---------------------------------------------------------
        def find_value(obj, wanted):
            if isinstance(obj, dict):
                for key, value in obj.items():
                    k = str(key).lower().replace("-", "_").replace(" ", "_")
                    if any(x in k for x in wanted):
                        if value is not None and not isinstance(value, (dict, list)):
                            return value

                for value in obj.values():
                    found = find_value(value, wanted)
                    if found not in (None, ""):
                        return found

            elif isinstance(obj, (list, tuple)):
                for value in obj:
                    found = find_value(value, wanted)
                    if found not in (None, ""):
                        return found

            return ""

        def side_value(row, side, field):
            if not isinstance(row, dict):
                return ""

            side = side.lower()

            # First try nested Excel/PDF/BOM objects.
            for key, value in row.items():
                k = str(key).lower().replace("-", "_").replace(" ", "_")

                if side in k and isinstance(value, dict):
                    result = find_value(value, field)
                    if result not in (None, ""):
                        return result

            # Then try explicit flat keys.
            candidates = []
            for f in field:
                candidates.extend([
                    f"{side}_{f}",
                    f"{f}_{side}",
                    f"{side}{f}",
                    f"{f}{side}",
                ])

            for candidate in candidates:
                for key, value in row.items():
                    k = str(key).lower().replace("-", "_").replace(" ", "_")
                    if k == candidate and value is not None:
                        return value

            return ""

        def clean(value):
            if value is None:
                return ""
            return str(value).strip()

        # ---------------------------------------------------------
        # Create workbook
        # ---------------------------------------------------------
        wb = Workbook()
        ws = wb.active
        ws.title = "BOM Comparison"

        headers = [
            "BOM1-PART",
            "BOM1-QTY",
            "BOM1-DESCRIPTION",
            "BOM1-RESULT",
            "QTY RESULT",
            "BOM2-PART",
            "BOM2-QTY",
            "BOM2-DESCRIPTION",
            "BOM2-RESULT",
        ]

        ws.append(headers)

        # ---------------------------------------------------------
        # Header formatting
        # ---------------------------------------------------------
        yellow = "FFFF00"
        green = "00B050"
        purple = "7030A0"

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

        for col in range(1, 4):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=yellow)

        for col in (4, 5, 9):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=green)

        for col in range(6, 9):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=purple)

        # ---------------------------------------------------------
        # Write all actual comparison results
        # ---------------------------------------------------------
        for row in results:
            excel_part = clean(side_value(row, "excel", ["part", "part_number", "item"]))
            excel_qty = clean(side_value(row, "excel", ["qty", "quantity"]))
            excel_desc = clean(side_value(row, "excel", ["description", "desc"]))

            pdf_part = clean(side_value(row, "pdf", ["part", "part_number", "item"]))
            pdf_qty = clean(side_value(row, "pdf", ["qty", "quantity"]))
            pdf_desc = clean(side_value(row, "pdf", ["description", "desc"]))

            # Fallback for result structures using BOM1/BOM2.
            if not excel_part:
                excel_part = clean(side_value(row, "bom1", ["part", "part_number", "item"]))
            if not excel_qty:
                excel_qty = clean(side_value(row, "bom1", ["qty", "quantity"]))
            if not excel_desc:
                excel_desc = clean(side_value(row, "bom1", ["description", "desc"]))

            if not pdf_part:
                pdf_part = clean(side_value(row, "bom2", ["part", "part_number", "item"]))
            if not pdf_qty:
                pdf_qty = clean(side_value(row, "bom2", ["qty", "quantity"]))
            if not pdf_desc:
                pdf_desc = clean(side_value(row, "bom2", ["description", "desc"]))

            status = clean(find_value(row, ["status", "result", "match_status"])).upper()

            if bom1_type == "Excel":
                b1_part, b1_qty, b1_desc = excel_part, excel_qty, excel_desc
                b2_part, b2_qty, b2_desc = pdf_part, pdf_qty, pdf_desc
            else:
                b1_part, b1_qty, b1_desc = pdf_part, pdf_qty, pdf_desc
                b2_part, b2_qty, b2_desc = excel_part, excel_qty, excel_desc

            # Use the actual verification status where available.
            if "MATCH" in status and "MISMATCH" not in status:
                b1_result = "EXISTS IN BOM2"
                b2_result = "EXISTS IN BOM1"
            elif "EXCEL_ONLY" in status:
                b1_result = "EXISTS IN BOM2" if bom1_type == "PDF" else "DOES NOT EXIST IN BOM2"
                b2_result = "DOES NOT EXIST IN BOM1" if bom1_type == "PDF" else "EXISTS IN BOM1"
            elif "PDF_ONLY" in status:
                b1_result = "EXISTS IN BOM2" if bom1_type == "Excel" else "DOES NOT EXIST IN BOM2"
                b2_result = "DOES NOT EXIST IN BOM1" if bom1_type == "Excel" else "EXISTS IN BOM1"
            else:
                b1_result = "EXISTS IN BOM2" if b1_part and b2_part else "DOES NOT EXIST IN BOM2"
                b2_result = "EXISTS IN BOM1" if b1_part and b2_part else "DOES NOT EXIST IN BOM1"

            if b1_part and b2_part:
                qty_result = (
                    "QUANTITIES MATCH"
                    if b1_qty == b2_qty
                    else "QUANTITIES DON'T MATCH"
                )
            else:
                qty_result = "N/A"

            ws.append([
                b1_part,
                b1_qty,
                b1_desc,
                b1_result,
                qty_result,
                b2_part,
                b2_qty,
                b2_desc,
                b2_result,
            ])

        # ---------------------------------------------------------
        # Borders / widths / alignment
        # ---------------------------------------------------------
        widths = {
            "A": 24,
            "B": 12,
            "C": 42,
            "D": 24,
            "E": 24,
            "F": 24,
            "G": 12,
            "H": 42,
            "I": 24,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(
                    left=Side(style="thin"),
                    right=Side(style="thin"),
                    top=Side(style="thin"),
                    bottom=Side(style="thin"),
                )
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # ---------------------------------------------------------
        # Save
        # ---------------------------------------------------------
        wb.save(output_path)
        return output_path

