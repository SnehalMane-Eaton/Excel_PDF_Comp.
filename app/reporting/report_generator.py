from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.matching.normalizer import normalize_part_number

HEADERS = [
    "BOM1-PART",
    "BOM1-QTY",
    "BOM1-DESCRIPTION",
    "BOM1-RESULT",
    "QTY RESULT",
    "BOM2-PART",
    "BOM2-QTY",
    "BOM2-DESCRIPTION",
    "BOM2-RESULT",
]

SUMMARY_SHEET = "SUMMARY"
AUDIT_SHEET = "VERIFICATION_AUDIT"
HEADER_YELLOW = "FFFF00"
HEADER_GREEN = "00B050"
HEADER_PURPLE = "7030A0"
RESULT_GREEN = "C6EFCE"
RESULT_RED = "FFC7CE"
RESULT_YELLOW = "FFEB9C"
RESULT_REVIEW = "FCE4D6"
BLACK = "000000"
WHITE = "FFFFFF"
BORDER = Border(
    left=Side(style="thin", color=BLACK),
    right=Side(style="thin", color=BLACK),
    top=Side(style="thin", color=BLACK),
    bottom=Side(style="thin", color=BLACK),
)


def generate_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    bom1_type, bom2_type = _normalize_bom_order(bom1_type, bom2_type)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)

    pair_results = _extract_pairs(verification_data)
    summary_sheet = workbook.create_sheet(SUMMARY_SHEET)
    _write_summary_sheet(summary_sheet, verification_data, pair_results, bom1_type, bom2_type)

    audit_sheet = workbook.create_sheet(AUDIT_SHEET)
    _write_audit_sheet(audit_sheet, pair_results)

    comparison_count = 0

    for index, pair_result in enumerate(pair_results, start=1):
        if not _should_create_pair_sheet(pair_result):
            continue

        comparison_count += 1
        comparison_sheet = workbook.create_sheet(
            _unique_sheet_name(workbook, _base_sheet_title(pair_result, index))
        )
        helper_sheet = workbook.create_sheet(
            _unique_sheet_name(workbook, f"_PAIR_{index:03d}_AUDIT")
        )
        helper_sheet.sheet_state = "hidden"

        _write_pair_sheet(
            comparison_sheet,
            helper_sheet,
            pair_result,
            bom1_type=bom1_type,
            bom2_type=bom2_type,
        )

    if comparison_count == 0:
        summary_sheet["A14"] = "No comparison sheets were created because no paired or partially extracted BOM data was available."

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    workbook.save(output)
    return str(output)


class ReportGenerator:
    def generate(
        self,
        verification_data: dict[str, Any],
        output_path: str | Path,
        bom1_type: str = "Excel",
        bom2_type: str = "PDF",
        **_: Any,
    ) -> str:
        return generate_report(verification_data, output_path, bom1_type, bom2_type)


def create_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    return generate_report(verification_data, output_path, bom1_type, bom2_type)


def save_report(
    verification_data: dict[str, Any],
    output_path: str | Path,
    bom1_type: str = "Excel",
    bom2_type: str = "PDF",
) -> str:
    return generate_report(verification_data, output_path, bom1_type, bom2_type)


def write_json(summary: Any, output_path: str | Path) -> str:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    if isinstance(summary, dict):
        payload = summary
    else:
        payload = _serialize_object(summary)

    output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return str(output)


def _normalize_bom_order(bom1_type: str, bom2_type: str) -> tuple[str, str]:
    bom1 = str(bom1_type or "").strip().title()
    bom2 = str(bom2_type or "").strip().title()

    if bom1 not in {"Excel", "Pdf", "PDF"}:
        raise ValueError("BOM1 must be either Excel or PDF.")
    if bom2 not in {"Excel", "Pdf", "PDF"}:
        raise ValueError("BOM2 must be either Excel or PDF.")

    bom1 = "PDF" if bom1 == "Pdf" else bom1
    bom2 = "PDF" if bom2 == "Pdf" else bom2

    if bom1 == bom2:
        raise ValueError("BOM1 and BOM2 must be different.")

    return bom1, bom2


def _extract_pairs(verification_data: dict[str, Any]) -> list[dict[str, Any]]:
    if isinstance(verification_data, dict) and isinstance(verification_data.get("pairs"), list):
        return verification_data["pairs"]

    if isinstance(verification_data, dict):
        return [verification_data]

    return []


def _write_summary_sheet(
    worksheet,
    verification_data: dict[str, Any],
    pair_results: list[dict[str, Any]],
    bom1_type: str,
    bom2_type: str,
) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = "Excel PDF BOM Verifier Summary"
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A2"] = (
        "Processing remains local. Select only company-approved non-synced folders. "
        "The application does not attempt to detect every sync provider."
    )
    worksheet["A3"] = (
        "PDF support is limited to the currently implemented structured BOM and engineering drawing formats; "
        "scanned/image-only PDFs are not reconstructed here."
    )
    worksheet["A4"] = f"Displayed order: BOM1={bom1_type}, BOM2={bom2_type}"

    summary = verification_data.get("summary", {}) if isinstance(verification_data, dict) else {}
    metrics = [
        ("Total file pairs", summary.get("total_pairs", len(pair_results))),
        ("Verified pairs", summary.get("verified", 0)),
        ("File processing errors", summary.get("processing_error", 0)),
        ("Excel-only files", summary.get("excel_only", 0)),
        ("PDF-only files", summary.get("pdf_only", 0)),
        ("Part MATCH", summary.get("match", 0)),
        ("Part QTY_MISMATCH", summary.get("qty_mismatch", 0)),
        ("Part EXCEL_ONLY", summary.get("bom_excel_only", 0)),
        ("Part PDF_ONLY", summary.get("bom_pdf_only", 0)),
        ("Part REVIEW", summary.get("bom_review", 0)),
    ]

    row = 6
    for label, value in metrics:
        worksheet.cell(row=row, column=1, value=label)
        worksheet.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    headers = [
        "Pair Key",
        "Status",
        "Message",
        "Excel File",
        "PDF File",
        "Warnings",
        "Comparison Sheet",
    ]
    _write_table_header(worksheet, row, headers, HEADER_GREEN, WHITE)
    row += 1

    if not pair_results:
        worksheet.cell(row=row, column=1, value="No file pairs were found.")
    else:
        for index, pair_result in enumerate(pair_results, start=1):
            comparison_name = _base_sheet_title(pair_result, index) if _should_create_pair_sheet(pair_result) else ""
            values = [
                pair_result.get("key", ""),
                pair_result.get("status", ""),
                pair_result.get("message", ""),
                pair_result.get("excel_file", ""),
                pair_result.get("pdf_file", ""),
                " | ".join(pair_result.get("warnings", [])),
                comparison_name,
            ]
            for column, value in enumerate(values, start=1):
                _write_plain_cell(worksheet, row, column, value)
            row += 1

    worksheet.auto_filter.ref = f"A{row - max(len(pair_results), 1)}:G{max(row, 1)}"
    _set_widths(
        worksheet,
        {"A": 24, "B": 18, "C": 42, "D": 38, "E": 38, "F": 42, "G": 24},
    )


def _write_audit_sheet(worksheet, pair_results: list[dict[str, Any]]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    headers = [
        "Pair Key",
        "Normalized Part",
        "Status",
        "Excel Quantity",
        "PDF Quantity",
        "Difference",
        "Description",
        "Excel Rows",
        "PDF Pages",
        "Remark",
    ]
    _write_table_header(worksheet, 1, headers, HEADER_GREEN, WHITE)

    row = 2
    for pair_result in pair_results:
        verification = pair_result.get("verification") or {}
        for result in verification.get("results", []):
            values = [
                pair_result.get("key", ""),
                result.get("part_number", ""),
                result.get("status", ""),
                result.get("excel_quantity", ""),
                result.get("pdf_quantity", ""),
                result.get("difference", ""),
                result.get("description", ""),
                ", ".join(str(value) for value in result.get("excel_rows", [])),
                ", ".join(str(value) for value in result.get("pdf_pages", [])),
                result.get("remark", ""),
            ]
            for column, value in enumerate(values, start=1):
                _write_plain_cell(worksheet, row, column, value)
            row += 1

    if row == 2:
        worksheet["A2"] = "No verification results were available."

    _set_widths(
        worksheet,
        {"A": 24, "B": 22, "C": 16, "D": 14, "E": 14, "F": 12, "G": 34, "H": 16, "I": 16, "J": 42},
    )


def _should_create_pair_sheet(pair_result: dict[str, Any]) -> bool:
    if pair_result.get("excel_records") or pair_result.get("pdf_records"):
        return True

    return pair_result.get("status") in {"VERIFIED", "PROCESSING_ERROR", "PAIRED"}


def _base_sheet_title(pair_result: dict[str, Any], index: int) -> str:
    key = str(pair_result.get("key") or "comparison").strip() or "comparison"
    safe = "".join(character if character not in '[]:*?/\\' else "_" for character in key)
    safe = safe[:22] if safe else "comparison"
    return f"{index:02d}_{safe}"


def _write_pair_sheet(
    worksheet,
    helper_sheet,
    pair_result: dict[str, Any],
    bom1_type: str,
    bom2_type: str,
) -> None:
    bom1_source = "excel" if bom1_type == "Excel" else "pdf"
    bom2_source = "pdf" if bom1_source == "excel" else "excel"
    bom1_records = _records_for_source(pair_result, bom1_source)
    bom2_records = _records_for_source(pair_result, bom2_source)
    row_count = max(len(bom1_records), len(bom2_records), 1)
    row_audits = _build_row_audits(pair_result, bom1_records, bom2_records, bom1_source, bom2_source, row_count)

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:I{row_count + 1}"
    _write_comparison_headers(worksheet)
    _write_helper_sheet(helper_sheet, row_audits)

    metadata = (
        f"Pair key: {pair_result.get('key', '')}\n"
        f"Status: {pair_result.get('status', '')}\n"
        f"Excel file: {pair_result.get('excel_file', '')}\n"
        f"PDF file: {pair_result.get('pdf_file', '')}\n"
        f"Displayed order: BOM1={bom1_type}, BOM2={bom2_type}\n"
        "BOM1/BOM2 rows are independent extracted source lists. Verification status remains authoritative on the audit sheet."
    )
    worksheet["A1"].comment = Comment(metadata, "Excel PDF BOM Verifier")

    for offset in range(row_count):
        row = offset + 2
        bom1_record = bom1_records[offset] if offset < len(bom1_records) else None
        bom2_record = bom2_records[offset] if offset < len(bom2_records) else None
        audit = row_audits[offset]

        _write_source_row(worksheet, row, 1, bom1_record)
        _write_source_row(worksheet, row, 6, bom2_record)

        worksheet.cell(row=row, column=4, value=f"=IF(COUNTA(A{row}:C{row})=0,\"\",'{helper_sheet.title}'!C{row})")
        worksheet.cell(row=row, column=5, value=f"=IF(AND(COUNTA(A{row}:C{row})=0,COUNTA(F{row}:H{row})=0),\"\",'{helper_sheet.title}'!D{row})")
        worksheet.cell(row=row, column=9, value=f"=IF(COUNTA(F{row}:H{row})=0,\"\",'{helper_sheet.title}'!E{row})")

        for column in (4, 5, 9):
            worksheet.cell(row=row, column=column).border = BORDER
            worksheet.cell(row=row, column=column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        _apply_result_fill(worksheet.cell(row=row, column=4), audit["bom1_result"])
        _apply_result_fill(worksheet.cell(row=row, column=5), audit["qty_result"])
        _apply_result_fill(worksheet.cell(row=row, column=9), audit["bom2_result"])
        _style_row(worksheet, row)

    _set_widths(
        worksheet,
        {"A": 20, "B": 12, "C": 38, "D": 24, "E": 24, "F": 20, "G": 12, "H": 38, "I": 24},
    )
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_area = f"A1:I{row_count + 1}"
    worksheet.print_title_rows = "1:1"


def _records_for_source(pair_result: dict[str, Any], source: str) -> list[dict[str, Any]]:
    key = "excel_records" if source == "excel" else "pdf_records"
    records = list(pair_result.get(key) or [])
    return records


def _build_row_audits(
    pair_result: dict[str, Any],
    bom1_records: list[dict[str, Any]],
    bom2_records: list[dict[str, Any]],
    bom1_source: str,
    bom2_source: str,
    row_count: int,
) -> list[dict[str, str]]:
    verification = pair_result.get("verification") or {}
    result_by_part: dict[str, dict[str, Any]] = {}
    blank_excel_reviews: dict[Any, dict[str, Any]] = {}

    for result in verification.get("results", []):
        part_number = result.get("part_number")
        if part_number:
            result_by_part[normalize_part_number(part_number)] = result
        elif result.get("status") == "REVIEW":
            for excel_row in result.get("excel_rows", []):
                blank_excel_reviews[excel_row] = result

    audits: list[dict[str, str]] = []

    for index in range(row_count):
        bom1_record = bom1_records[index] if index < len(bom1_records) else None
        bom2_record = bom2_records[index] if index < len(bom2_records) else None
        bom1_lookup = _lookup_verification_result(bom1_record, bom1_source, result_by_part, blank_excel_reviews)
        bom2_lookup = _lookup_verification_result(bom2_record, bom2_source, result_by_part, blank_excel_reviews)
        qty_result = _qty_result_text(bom1_lookup or bom2_lookup, bom1_record or bom2_record)

        audits.append(
            {
                "bom1_normalized": _normalized_part(bom1_record),
                "bom2_normalized": _normalized_part(bom2_record),
                "bom1_result": _existence_text(bom1_lookup, bom1_source, bom1_record),
                "qty_result": qty_result,
                "bom2_result": _existence_text(bom2_lookup, bom2_source, bom2_record),
            }
        )

    return audits


def _lookup_verification_result(
    record: dict[str, Any] | None,
    source: str,
    result_by_part: dict[str, dict[str, Any]],
    blank_excel_reviews: dict[Any, dict[str, Any]],
) -> dict[str, Any] | None:
    if not record:
        return None

    normalized = _normalized_part(record)
    if normalized:
        return result_by_part.get(normalized)

    if source == "excel":
        return blank_excel_reviews.get(record.get("excel_row"))

    return None


def _normalized_part(record: dict[str, Any] | None) -> str:
    if not record:
        return ""
    return normalize_part_number(record.get("part_number"))


def _existence_text(
    result: dict[str, Any] | None,
    source: str,
    record: dict[str, Any] | None,
) -> str:
    if not record:
        return ""

    if source == "excel":
        present = _has_other_side_presence(result, "pdf")
        return "EXISTS IN BOM2" if present else "DOES NOT EXIST IN BOM2"

    present = _has_other_side_presence(result, "excel")
    return "EXISTS IN BOM1" if present else "DOES NOT EXIST IN BOM1"


def _has_other_side_presence(result: dict[str, Any] | None, source: str) -> bool:
    if not result:
        return False

    status = result.get("status")
    if status in {"MATCH", "QTY_MISMATCH"}:
        return True

    if source == "pdf":
        return status == "REVIEW" and bool(result.get("pdf_pages"))

    return status == "REVIEW" and bool(result.get("excel_rows"))


def _qty_result_text(
    result: dict[str, Any] | None,
    record: dict[str, Any] | None,
) -> str:
    if not record:
        return ""

    if not result:
        return "N/A"

    status = result.get("status")
    if status == "MATCH":
        return "QUANTITIES MATCH"
    if status == "QTY_MISMATCH":
        return "QUANTITIES DON'T MATCH"
    if status == "REVIEW":
        return "REVIEW"
    return "N/A"


def _write_helper_sheet(helper_sheet, row_audits: list[dict[str, str]]) -> None:
    helper_sheet.append(["bom1_normalized", "bom2_normalized", "bom1_result", "qty_result", "bom2_result"])
    for audit in row_audits:
        helper_sheet.append(
            [
                audit["bom1_normalized"],
                audit["bom2_normalized"],
                audit["bom1_result"],
                audit["qty_result"],
                audit["bom2_result"],
            ]
        )


def _write_comparison_headers(worksheet) -> None:
    for column, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if column in (1, 2, 3):
            cell.fill = PatternFill("solid", fgColor=HEADER_YELLOW)
            cell.font = Font(bold=True, color=BLACK)
        elif column in (4, 5, 9):
            cell.fill = PatternFill("solid", fgColor=HEADER_GREEN)
            cell.font = Font(bold=True, color=WHITE)
        else:
            cell.fill = PatternFill("solid", fgColor=HEADER_PURPLE)
            cell.font = Font(bold=True, color=WHITE)
    worksheet.row_dimensions[1].height = 28


def _write_source_row(worksheet, row: int, start_column: int, record: dict[str, Any] | None) -> None:
    part = record.get("part_number") if record else ""
    quantity = record.get("quantity") if record else ""
    description = record.get("description") if record else ""

    _write_text_cell(worksheet, row, start_column, part)
    _write_quantity_cell(worksheet, row, start_column + 1, quantity)
    _write_text_cell(worksheet, row, start_column + 2, description)


def _write_text_cell(worksheet, row: int, column: int, value: Any) -> None:
    text = "" if value is None else str(value)
    cell = worksheet.cell(row=row, column=column, value=text)
    cell.data_type = "s"
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    cell.font = Font(name="Calibri", size=11, color=BLACK)


def _write_quantity_cell(worksheet, row: int, column: int, value: Any) -> None:
    cell = worksheet.cell(row=row, column=column)
    if value is None or value == "":
        cell.value = ""
        cell.data_type = "s"
    elif isinstance(value, bool):
        cell.value = str(value)
        cell.data_type = "s"
    else:
        try:
            number = float(value)
            cell.value = int(number) if number.is_integer() else number
        except (TypeError, ValueError):
            cell.value = str(value)
            cell.data_type = "s"
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    cell.font = Font(name="Calibri", size=11, color=BLACK)


def _style_row(worksheet, row: int) -> None:
    max_length = 0
    for column in (3, 8):
        value = worksheet.cell(row=row, column=column).value or ""
        max_length = max(max_length, len(str(value)))

    worksheet.row_dimensions[row].height = max(22, min(90, 18 + (max_length // 45) * 14))


def _apply_result_fill(cell, result_text: str) -> None:
    if result_text in {"EXISTS IN BOM2", "EXISTS IN BOM1", "QUANTITIES MATCH"}:
        color = RESULT_GREEN
    elif result_text in {"DOES NOT EXIST IN BOM2", "DOES NOT EXIST IN BOM1", "QUANTITIES DON'T MATCH"}:
        color = RESULT_RED
    elif result_text == "REVIEW":
        color = RESULT_REVIEW
    else:
        color = RESULT_YELLOW

    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name="Calibri", size=11, color=BLACK)


def _write_table_header(worksheet, row: int, headers: list[str], fill: str, font_color: str) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=font_color)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_plain_cell(worksheet, row: int, column: int, value: Any) -> None:
    cell = worksheet.cell(row=row, column=column, value="" if value is None else value)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _set_widths(worksheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _unique_sheet_name(workbook: Workbook, base_name: str) -> str:
    candidate = base_name[:31]
    if candidate not in workbook.sheetnames:
        return candidate

    index = 2
    while True:
        suffix = f"_{index}"
        candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
        index += 1


def _serialize_object(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, list):
        return [_serialize_object(item) for item in value]
    if isinstance(value, tuple):
        return [_serialize_object(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _serialize_object(item) for key, item in value.items()}
    if hasattr(value, "__dict__"):
        return _serialize_object(vars(value))
    return str(value)
