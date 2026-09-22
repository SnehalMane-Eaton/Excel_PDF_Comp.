from pathlib import Path

p = Path("app/reporting/report_generator.py")
s = p.read_text(encoding="utf-8")

# Replace the compatibility wrapper only.
start = s.find("class ReportGenerator:")
if start == -1:
    raise SystemExit("ReportGenerator wrapper not found.")

new_wrapper = r'''
class ReportGenerator:
    """
    UI compatibility wrapper.

    The verification engine returns:
        batch_result["pairs"][...]["verification"]["results"]

    The report is generated directly from those actual comparison
    results so the selected Excel/PDF BOM data appears in Excel.
    """

    def generate(
        self,
        verification_data,
        output_path,
        bom1_type="Excel",
        bom2_type="PDF",
        **kwargs,
    ):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.utils import get_column_letter

        # ---------------------------------------------------------
        # Collect actual verification results from BatchProcessor
        # ---------------------------------------------------------
        results = []

        if isinstance(verification_data, dict):
            for pair in verification_data.get("pairs", []):
                if not isinstance(pair, dict):
                    continue

                verification = pair.get("verification", {})
                if isinstance(verification, dict):
                    rows = verification.get("results", [])
                    if isinstance(rows, list):
                        results.extend(rows)

        # ---------------------------------------------------------
        # Helpers for different result dictionary structures
        # ---------------------------------------------------------
        def find_value(obj, wanted):
            if isinstance(obj, dict):
                for key, value in obj.items():
                    k = str(key).lower().replace("-", "_").replace(" ", "_")
                    if any(x in k for x in wanted):
                        if value is not None and not isinstance(value, (dict, list)):
                            return value

                for value in obj.values():
                    found = find_value(value, wanted)
                    if found not in (None, ""):
                        return found

            elif isinstance(obj, (list, tuple)):
                for value in obj:
                    found = find_value(value, wanted)
                    if found not in (None, ""):
                        return found

            return ""

        def side_value(row, side, field):
            if not isinstance(row, dict):
                return ""

            side = side.lower()

            # First try nested Excel/PDF/BOM objects.
            for key, value in row.items():
                k = str(key).lower().replace("-", "_").replace(" ", "_")

                if side in k and isinstance(value, dict):
                    result = find_value(value, field)
                    if result not in (None, ""):
                        return result

            # Then try explicit flat keys.
            candidates = []
            for f in field:
                candidates.extend([
                    f"{side}_{f}",
                    f"{f}_{side}",
                    f"{side}{f}",
                    f"{f}{side}",
                ])

            for candidate in candidates:
                for key, value in row.items():
                    k = str(key).lower().replace("-", "_").replace(" ", "_")
                    if k == candidate and value is not None:
                        return value

            return ""

        def clean(value):
            if value is None:
                return ""
            return str(value).strip()

        # ---------------------------------------------------------
        # Create workbook
        # ---------------------------------------------------------
        wb = Workbook()
        ws = wb.active
        ws.title = "BOM Comparison"

        headers = [
            "BOM1-PART",
            "BOM1-QTY",
            "BOM1-DESCRIPTION",
            "BOM1-RESULT",
            "QTY RESULT",
            "BOM2-PART",
            "BOM2-QTY",
            "BOM2-DESCRIPTION",
            "BOM2-RESULT",
        ]

        ws.append(headers)

        # ---------------------------------------------------------
        # Header formatting
        # ---------------------------------------------------------
        yellow = "FFFF00"
        green = "00B050"
        purple = "7030A0"

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

        for col in range(1, 4):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=yellow)

        for col in (4, 5, 9):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=green)

        for col in range(6, 9):
            ws.cell(1, col).fill = PatternFill("solid", fgColor=purple)

        # ---------------------------------------------------------
        # Write all actual comparison results
        # ---------------------------------------------------------
        for row in results:
            excel_part = clean(side_value(row, "excel", ["part", "part_number", "item"]))
            excel_qty = clean(side_value(row, "excel", ["qty", "quantity"]))
            excel_desc = clean(side_value(row, "excel", ["description", "desc"]))

            pdf_part = clean(side_value(row, "pdf", ["part", "part_number", "item"]))
            pdf_qty = clean(side_value(row, "pdf", ["qty", "quantity"]))
            pdf_desc = clean(side_value(row, "pdf", ["description", "desc"]))

            # Fallback for result structures using BOM1/BOM2.
            if not excel_part:
                excel_part = clean(side_value(row, "bom1", ["part", "part_number", "item"]))
            if not excel_qty:
                excel_qty = clean(side_value(row, "bom1", ["qty", "quantity"]))
            if not excel_desc:
                excel_desc = clean(side_value(row, "bom1", ["description", "desc"]))

            if not pdf_part:
                pdf_part = clean(side_value(row, "bom2", ["part", "part_number", "item"]))
            if not pdf_qty:
                pdf_qty = clean(side_value(row, "bom2", ["qty", "quantity"]))
            if not pdf_desc:
                pdf_desc = clean(side_value(row, "bom2", ["description", "desc"]))

            status = clean(find_value(row, ["status", "result", "match_status"])).upper()

            if bom1_type == "Excel":
                b1_part, b1_qty, b1_desc = excel_part, excel_qty, excel_desc
                b2_part, b2_qty, b2_desc = pdf_part, pdf_qty, pdf_desc
            else:
                b1_part, b1_qty, b1_desc = pdf_part, pdf_qty, pdf_desc
                b2_part, b2_qty, b2_desc = excel_part, excel_qty, excel_desc

            # Use the actual verification status where available.
            if "MATCH" in status and "MISMATCH" not in status:
                b1_result = "EXISTS IN BOM2"
                b2_result = "EXISTS IN BOM1"
            elif "EXCEL_ONLY" in status:
                b1_result = "EXISTS IN BOM2" if bom1_type == "PDF" else "DOES NOT EXIST IN BOM2"
                b2_result = "DOES NOT EXIST IN BOM1" if bom1_type == "PDF" else "EXISTS IN BOM1"
            elif "PDF_ONLY" in status:
                b1_result = "EXISTS IN BOM2" if bom1_type == "Excel" else "DOES NOT EXIST IN BOM2"
                b2_result = "DOES NOT EXIST IN BOM1" if bom1_type == "Excel" else "EXISTS IN BOM1"
            else:
                b1_result = "EXISTS IN BOM2" if b1_part and b2_part else "DOES NOT EXIST IN BOM2"
                b2_result = "EXISTS IN BOM1" if b1_part and b2_part else "DOES NOT EXIST IN BOM1"

            if b1_part and b2_part:
                qty_result = (
                    "QUANTITIES MATCH"
                    if b1_qty == b2_qty
                    else "QUANTITIES DON'T MATCH"
                )
            else:
                qty_result = "N/A"

            ws.append([
                b1_part,
                b1_qty,
                b1_desc,
                b1_result,
                qty_result,
                b2_part,
                b2_qty,
                b2_desc,
                b2_result,
            ])

        # ---------------------------------------------------------
        # Borders / widths / alignment
        # ---------------------------------------------------------
        widths = {
            "A": 24,
            "B": 12,
            "C": 42,
            "D": 24,
            "E": 24,
            "F": 24,
            "G": 12,
            "H": 42,
            "I": 24,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(
                    left=Side(style="thin"),
                    right=Side(style="thin"),
                    top=Side(style="thin"),
                    bottom=Side(style="thin"),
                )
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # ---------------------------------------------------------
        # Save
        # ---------------------------------------------------------
        wb.save(output_path)
        return output_path
'''

p.write_text(s[:start] + new_wrapper + "\n", encoding="utf-8")

compile(p.read_text(encoding="utf-8"), str(p), "exec")
print("REPORT GENERATOR REPLACED SUCCESSFULLY")
